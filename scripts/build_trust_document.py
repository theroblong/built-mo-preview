"""Build Brian's trust document ("How Mo works") in Aevah branding: Word, verified PDF, editable HTML.

# protocol: not-a-forecast-arm

Source of the words: docs/drafts/aevah_trust_document_DRAFT.md (edit the words there, then re-run this).
Rules for the content: docs/drafts/aevah_trust_document_rob_checklist.md (Rob's answers, 2026-10-14).
Brand: docs/aevah_brand/brand/aevah-brand-guide.md (colors, Newsreader/Inter with Georgia/Arial fallbacks,
face-mark lockup, conceptual diagrams only, amber as the single highlight, never amber text on light).

Run from the repo root:  python3 scripts/build_trust_document.py [--no-pdf]
"""
from __future__ import annotations

import argparse
import base64
import html
import re
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt, RGBColor  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from MO_142_build_brief import to_pdf  # noqa: E402  (Word export with unique temp name + text check)

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs/drafts/aevah_trust_document_DRAFT.md"
BRAND = ROOT / "docs/aevah_brand/brand/assets/generated"
VERSION = "Draft v0.3"
ISSUED = "October 2026"
STEM = "Aevah_How_Mo_Works_v0.3_2026-10-14"
OUT = ROOT / "docs/briefs/trust_document_v0.3"

# Aevah palette (brand guide 8.2)
PLUM_DEEP, PLUM, VIOLET, INK, INK_SOFT, MUTED = "#170A2A", "#462453", "#77548E", "#29232E", "#514758", "#625A69"
PAPER, LILAC_SOFT, LILAC, RULE, AMBER, AMBER_SOFT = "#FFFDFA", "#F5F2F6", "#EFE8F2", "#E2D8E5", "#F5A65B", "#F9E9C9"
SERIF_DOCX, SANS_DOCX = "Georgia", "Arial"

FLOW = [  # conceptual only (brand 9.1): no product screens, schemas or mechanics; detail is in the numbered list
    "Source\nrecords",
    "Data store on\nBUILT's Azure\nsubscription",
    "Calculation",
    "Stored results\nwith model\nversion, date",
    "Mo app and\nMo Chat",
]


# ---------------------------------------------------------------- parse the draft
def parse(md: str) -> list[tuple]:
    """Blocks from the draft's small Markdown subset: #, ##, paragraphs, '- ' bullets, '1. ' steps, | tables."""
    md = re.sub(r"<!--.*?-->", "", md, flags=re.S).strip()
    B: list[tuple] = []
    for chunk in re.split(r"\n\s*\n", md):
        lines = chunk.strip("\n").split("\n")
        first = lines[0]
        if first.startswith("# "):
            B.append(("title", first[2:].strip()))
        elif first.startswith("### "):
            B.append(("h3", first[4:].strip()))
        elif first.startswith("## "):
            B.append(("h2", first[3:].strip()))
        elif first.startswith("!["):
            m = re.match(r"!\[(.*)\]\((.+)\)", " ".join(ln.strip() for ln in lines))
            B.append(("img", SRC.parent / m.group(2), m.group(1)))
        elif first.startswith("|"):
            rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines]
            B.append(("table", rows[0], [r for r in rows[2:]]))
        elif re.match(r"^(- |\d+\. )", first):
            items, kind = [], "steps" if first[0].isdigit() else "bullets"
            for ln in lines:
                if re.match(r"^(- |\d+\. )", ln):
                    items.append(re.sub(r"^(- |\d+\. )", "", ln).strip())
                else:
                    items[-1] += " " + ln.strip()
            B.append((kind, items))
        else:
            B.append(("p", " ".join(ln.strip() for ln in lines)))
    # title block: the bold line after the title is the subtitle, the next line is the byline
    assert B[0][0] == "title" and B[1][1].startswith("**") and B[2][0] == "p"
    B[1] = ("subtitle", B[1][1].strip("*"))
    B[2] = ("meta", B[2][1])
    return B


def runs(text: str) -> list[tuple[str, bool, bool]]:
    """Inline **bold** and *italic* into (text, bold, italic) runs."""
    out = []
    for part in re.split(r"(\*\*.+?\*\*|\*[^*]+?\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("*"):
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out


def plain(text: str) -> str:
    return "".join(t for t, _, _ in runs(text))


# ---------------------------------------------------------------- diagram
def flow_diagram(path: Path) -> Path:
    """Five porcelain tiles on lilac, violet connectors, one amber highlight on the outcome (brand 9, style A).
    Labels are about 14 px or larger when the image spans the 820 px HTML column (brand 8.3 minimum)."""
    plt.rcParams["font.family"] = ["Arial", "DejaVu Sans"]  # brand sans fallback (Avenir Next .ttc loads bold only)
    fig, ax = plt.subplots(figsize=(6.9, 1.75), dpi=240)
    fig.patch.set_facecolor(PAPER)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 25)
    ax.axis("off")
    ax.add_patch(FancyBboxPatch((0.5, 0.5), 99, 24, boxstyle="round,pad=0,rounding_size=2.2",
                                fc="#F7F4FA", ec=RULE, lw=0.8))
    n, w, gap = len(FLOW), 17.0, 2.6
    x0 = (100 - (n * w + (n - 1) * gap)) / 2
    for i, head in enumerate(FLOW):
        x = x0 + i * (w + gap)
        last = i == n - 1
        if last:  # the single amber highlight: the outcome BUILT's teams use
            ax.add_patch(FancyBboxPatch((x - 0.7, 2.8), w + 1.4, 19.4, boxstyle="round,pad=0,rounding_size=2.4",
                                        fc=AMBER_SOFT, ec="none", alpha=0.9))
        ax.add_patch(FancyBboxPatch((x, 3.5), w, 18, boxstyle="round,pad=0,rounding_size=1.8",
                                    fc="#FFFFFF", ec=AMBER if last else RULE, lw=1.4 if last else 0.9))
        ax.text(x + 1.6, 20.0, f"{i + 1}", color=VIOLET, fontsize=8.5, fontweight="bold", va="top")
        ax.text(x + w / 2, 12.0, head, color=PLUM_DEEP, fontsize=8.6, fontweight="bold", ha="center", va="center",
                linespacing=1.2)
        if i < n - 1:
            ax.add_patch(FancyArrowPatch((x + w + 0.3, 12.5), (x + w + gap - 0.3, 12.5), arrowstyle="-|>",
                                         mutation_scale=9, color=VIOLET, lw=1.3))
    fig.savefig(path, facecolor=PAPER, bbox_inches="tight", pad_inches=0.02)
    plt.close(fig)
    return path


# ---------------------------------------------------------------- Word
def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#"))


def shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:color", "auto"), ("w:fill", hex_color.lstrip("#"))):
        sh.set(qn(k), v)
    tcPr.append(sh)


def borders(table, color: str = RULE, inside: bool = True) -> None:
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        on = inside or edge in ("top", "left", "bottom", "right")
        e.set(qn("w:val"), "single" if on else "nil")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), color.lstrip("#"))
        b.append(e)
    tblPr.append(b)


def cell_margins(table, twips: int = 100) -> None:
    m = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:w"), str(twips))
        e.set(qn("w:type"), "dxa")
        m.append(e)
    table._tbl.tblPr.append(m)


def add_runs(p, text: str, size: float | None = None, color: str | None = None) -> None:
    for t, bold, italic in runs(text):
        r = p.add_run(t)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if bold:
            r.font.color.rgb = rgb(PLUM_DEEP)
        elif italic:
            r.font.name, r.font.color.rgb = SERIF_DOCX, rgb(VIOLET)
        elif color:
            r.font.color.rgb = rgb(color)


def render_docx(B: list[tuple], diagram: Path, path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin, sec.bottom_margin = Inches(0.7), Inches(0.75)
    width = sec.page_width - sec.left_margin - sec.right_margin
    st = doc.styles["Normal"]
    st.font.name, st.font.size, st.font.color.rgb = SANS_DOCX, Pt(10.5), rgb(INK)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), SANS_DOCX)
    st.paragraph_format.space_after, st.paragraph_format.line_spacing = Pt(7), 1.25
    h = doc.styles["Heading 2"]
    h.font.name, h.font.size, h.font.color.rgb, h.font.bold = SERIF_DOCX, Pt(17), rgb(PLUM), False
    h.element.rPr.rFonts.set(qn("w:asciiTheme"), "")  # use Georgia, not the theme heading font
    h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(18), Pt(6)
    lb = doc.styles["List Bullet"]
    lb.font.name, lb.font.size = SANS_DOCX, Pt(10.5)
    foot = sec.footer.paragraphs[0]
    foot.text = f"How Mo works  ·  Aevah for BUILT  ·  {ISSUED}  ·  {VERSION}"
    foot.runs[0].font.size, foot.runs[0].font.color.rgb = Pt(8), rgb(MUTED)

    section = ""
    for b in B:
        k = b[0]
        if k == "title":
            doc.add_picture(str(BRAND / "logo-lockup-light-transparent.png"), width=Inches(1.55))
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            r = p.add_run(b[1])
            r.font.name, r.font.size, r.font.color.rgb = SANS_DOCX, Pt(32), rgb(PLUM_DEEP)
            p.paragraph_format.space_after = Pt(4)
        elif k == "subtitle":
            p = doc.add_paragraph()
            add_runs(p, b[1], size=13.5, color=INK_SOFT)
        elif k == "meta":
            p = doc.add_paragraph()
            add_runs(p, b[1], size=9, color=MUTED)
            _rule_below(p)
        elif k == "h2":
            section = b[1]
            doc.add_heading(b[1], level=2).paragraph_format.keep_with_next = True
        elif k == "h3":
            p = doc.add_paragraph()
            p.paragraph_format.space_before, p.paragraph_format.keep_with_next = Pt(10), True
            r = p.add_run(b[1])
            r.font.name, r.font.size, r.font.color.rgb = SERIF_DOCX, Pt(13), rgb(VIOLET)
        elif k == "p":
            p = doc.add_paragraph()
            if b[1].startswith("Mo runs on BUILT's own Azure"):  # the one security sentence, set as a panel
                t = doc.add_table(rows=1, cols=1)
                borders(t, LILAC)
                cell_margins(t, 160)
                c = t.cell(0, 0)
                shade(c, LILAC_SOFT)
                add_runs(c.paragraphs[0], b[1], size=11.5, color=PLUM_DEEP)
                p._p.getparent().remove(p._p)
                doc.add_paragraph().paragraph_format.space_after = Pt(0)
                continue
            if "[[" in b[1]:
                r = p.add_run(b[1])
                r.font.highlight_color = 7  # yellow: placeholder to fill before sharing
                continue
            add_runs(p, b[1])
        elif k == "img":
            doc.add_picture(str(b[1]), width=width)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
            p = doc.add_paragraph()
            r = p.add_run(b[2])
            r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), rgb(MUTED)
        elif k == "bullets" and section == "In brief":  # three cards
            t = doc.add_table(rows=1, cols=len(b[1]))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(t, PAPER)
            cell_margins(t, 140)
            for j, item in enumerate(b[1]):
                c = t.cell(0, j)
                shade(c, LILAC_SOFT)
                head, _, rest = item.partition("** ")
                p = c.paragraphs[0]
                r = p.add_run(head.strip("*"))
                r.bold, r.font.size, r.font.color.rgb = True, Pt(10.5), rgb(PLUM)
                p.paragraph_format.space_after = Pt(4)
                p2 = c.add_paragraph()
                add_runs(p2, rest, size=9.5, color=INK_SOFT)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k == "bullets":
            for item in b[1]:
                add_runs(doc.add_paragraph(style="List Bullet"), item)
        elif k == "steps":
            doc.add_picture(str(diagram), width=width)
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            t = doc.add_table(rows=len(b[1]), cols=2)
            borders(t, RULE, inside=False)
            cell_margins(t, 90)
            for i, item in enumerate(b[1]):
                n_cell, c = t.cell(i, 0), t.cell(i, 1)
                n_cell.width, c.width = Inches(0.4), width - Inches(0.4)
                r = n_cell.paragraphs[0].add_run(str(i + 1))
                r.bold, r.font.color.rgb, r.font.size = True, rgb(VIOLET), Pt(11)
                add_runs(c.paragraphs[0], item, size=10)
                for cc in (n_cell, c):
                    cc.paragraphs[0].paragraph_format.space_after = Pt(0)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif k == "table":
            head, rows = b[1], b[2]
            t = doc.add_table(rows=1 + len(rows), cols=len(head))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(t, RULE)
            cell_margins(t, 90)
            for j, x in enumerate(head):
                c = t.cell(0, j)
                shade(c, PLUM_DEEP)
                r = c.paragraphs[0].add_run(plain(x))
                r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), rgb("#FFFFFF")
            trPr = t.rows[0]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:tblHeader"))
            for i, row in enumerate(rows, start=1):
                t.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for j, v in enumerate(row):
                    c = t.cell(i, j)
                    if i % 2 == 0:
                        shade(c, LILAC_SOFT)
                    p = c.paragraphs[0]
                    p.paragraph_format.space_after = Pt(0)
                    if j == 0:
                        r = p.add_run(plain(v))
                        r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), rgb(PLUM_DEEP)
                    else:
                        add_runs(p, v, size=9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
    doc.core_properties.title = "How Mo works"
    doc.core_properties.author = "Aevah"
    doc.save(path)


def _spacing(twentieths: int):
    e = OxmlElement("w:spacing")
    e.set(qn("w:val"), str(twentieths))
    return e


def _rule_below(p) -> None:
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "8"), ("w:color", RULE.lstrip("#"))):
        bot.set(qn(k), v)
    bdr.append(bot)
    pPr.append(bdr)
    p.paragraph_format.space_after = Pt(10)


# ---------------------------------------------------------------- HTML (editable, single light look like the print)
CSS = f"""
:root {{ --paper:{PAPER}; --plum-deep:{PLUM_DEEP}; --plum:{PLUM}; --violet:{VIOLET}; --ink:{INK}; --ink-soft:{INK_SOFT};
  --muted:{MUTED}; --lilac-soft:{LILAC_SOFT}; --lilac:{LILAC}; --rule:{RULE}; --amber:{AMBER}; --amber-soft:{AMBER_SOFT};
  color-scheme: light; }}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--paper); color:var(--ink);
  font:16px/1.7 Inter, Avenir, "Avenir Next", "Segoe UI", Helvetica, Arial, sans-serif; }}
.bar {{ position:sticky; top:0; z-index:5; background:var(--plum-deep); color:#fff; display:flex; flex-wrap:wrap; gap:10px;
  align-items:center; padding:10px 16px; }}
.bar b {{ margin-right:auto; font-weight:600; }}
.bar button {{ background:#fff; color:var(--plum-deep); border:0; border-radius:7px; padding:7px 12px;
  font:600 13px Inter, Arial, sans-serif; cursor:pointer; }}
.bar button.primary {{ background:var(--amber); color:var(--plum-deep); }}
.bar .hint {{ font-size:12px; opacity:.85; width:100%; }}
main {{ max-width:820px; margin:0 auto; padding:40px 16px 72px; }}
main[contenteditable=true] {{ outline:none; }}
main[contenteditable=true] :is(p,li,td,th,h1,h2):hover {{ background:var(--lilac-soft); }}
.logo {{ width:150px; height:auto; display:block; margin:0 0 28px -6px; }}
.eyebrow {{ color:var(--violet); font-weight:700; font-size:12px; letter-spacing:.13em; margin:0 0 6px; }}
h1 {{ font-weight:540; font-size:clamp(34px,6vw,52px); letter-spacing:-.05em; line-height:1.06; color:var(--plum-deep);
  margin:0 0 12px; text-wrap:balance; }}
.subtitle {{ color:var(--ink-soft); font-size:19px; line-height:1.5; margin:0 0 10px; max-width:60ch; }}
.meta {{ color:var(--muted); font-size:14px; margin:0 0 8px; padding-bottom:22px; border-bottom:1px solid var(--rule); }}
h2 {{ font-family:Newsreader, Baskerville, Georgia, serif; font-weight:470; font-size:30px; letter-spacing:-.035em;
  line-height:1.1; color:var(--plum); margin:48px 0 12px; text-wrap:balance; }}
p, li {{ max-width:68ch; }}
p strong, li strong, td strong {{ color:var(--plum-deep); font-weight:650; }}
em {{ font-family:Newsreader, Baskerville, Georgia, serif; color:var(--violet); font-size:1.08em; }}
li {{ margin:6px 0; }}
.cards {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(220px,1fr)); gap:14px; margin:18px 0 6px; }}
.card {{ background:var(--lilac-soft); border:1px solid var(--rule); border-radius:12px; padding:16px 18px; min-width:0; }}
h3.sub {{ font-family:Newsreader, Baskerville, Georgia, serif; font-weight:470; font-size:22px; color:var(--violet);
  letter-spacing:-.02em; margin:28px 0 6px; }}
.card h3 {{ font-size:16px; font-weight:600; line-height:1.3; color:var(--plum); margin:0 0 6px; }}
.card p {{ margin:0; font-size:15px; color:var(--ink-soft); }}
figure {{ margin:18px 0 10px; }}
figure img {{ width:100%; height:auto; display:block; }}
figure.wide img {{ border:1px solid var(--rule); border-radius:12px; }}
figcaption {{ color:var(--muted); font-size:14px; font-style:italic; margin-top:6px; }}
ol.steps {{ list-style:none; padding:0; margin:8px 0 0; counter-reset:s; }}
ol.steps li {{ counter-increment:s; display:grid; grid-template-columns:34px 1fr; gap:6px; padding:10px 0;
  border-bottom:1px solid var(--rule); margin:0; max-width:none; }}
ol.steps li::before {{ content:counter(s); color:var(--violet); font-weight:700; }}
.tw {{ overflow-x:auto; margin:14px 0 8px; }}
table {{ border-collapse:collapse; width:100%; font-size:14.5px; line-height:1.5; }}
th {{ background:var(--plum-deep); color:#fff; text-align:left; font-weight:600; padding:9px 12px; vertical-align:bottom; }}
td {{ border-bottom:1px solid var(--rule); padding:9px 12px; vertical-align:top; }}
td:first-child {{ font-weight:600; color:var(--plum-deep); }}
tr:nth-child(even) td {{ background:var(--lilac-soft); }}
.panel {{ background:var(--lilac-soft); border:1px solid var(--lilac); border-radius:14px; padding:18px 20px; font-size:18px;
  color:var(--plum-deep); max-width:none; }}
.todo {{ background:#FFF3B0; }}
footer {{ color:var(--muted); font-size:13px; margin-top:56px; padding-top:14px; border-top:1px solid var(--rule); }}
@media print {{
  @page {{ size: letter; margin: 0.6in; }}
  body {{ font-size:10.5pt; }} .bar {{ display:none; }} main {{ max-width:none; padding:0; }}
  h2 {{ break-after:avoid; }} figure, table, .cards, .panel {{ break-inside:avoid; }}
  main :hover {{ background:none !important; }}
}}
"""

JS = """
const main = document.querySelector('main');
const toggle = document.getElementById('edit');
function setEdit(on){ main.contentEditable = on ? 'true' : 'false'; toggle.textContent = on ? 'Editing: ON' : 'Editing: OFF'; }
toggle.onclick = () => setEdit(main.contentEditable !== 'true');
document.getElementById('save').onclick = () => {
  const wasOn = main.contentEditable === 'true'; setEdit(false);
  const src = '<!doctype html>\\n' + document.documentElement.outerHTML; setEdit(wasOn);
  const a = document.createElement('a');
  a.href = URL.createObjectURL(new Blob([src], {type:'text/html'}));
  a.download = 'Aevah_How_Mo_Works_edited.html'; document.body.appendChild(a); a.click(); a.remove();
};
document.getElementById('print').onclick = () => { setEdit(false); window.print(); };
setEdit(false);
"""


def h_inline(text: str) -> str:
    out = []
    for t, bold, italic in runs(text):
        t = html.escape(t)
        out.append(f"<strong>{t}</strong>" if bold else f"<em>{t}</em>" if italic else t)
    return "".join(out)


def b64(p: Path) -> str:
    return base64.b64encode(p.read_bytes()).decode()


def render_html(B: list[tuple], diagram: Path, path: Path) -> None:
    out, section = [], ""
    for b in B:
        k = b[0]
        if k == "title":
            out.append(f'<img class="logo" alt="Aevah" src="data:image/png;base64,{b64(BRAND / "logo-lockup-on-paper.png")}">'
                       f'<h1>{html.escape(b[1])}</h1>')
        elif k == "subtitle":
            out.append(f'<p class="subtitle">{h_inline(b[1])}</p>')
        elif k == "meta":
            out.append(f'<p class="meta">{h_inline(b[1])}</p>')
        elif k == "h2":
            section = b[1]
            out.append(f"<h2>{html.escape(b[1])}</h2>")
        elif k == "h3":
            out.append(f'<h3 class="sub">{html.escape(b[1])}</h3>')
        elif k == "p":
            cls = ' class="panel"' if b[1].startswith("Mo runs on BUILT's own Azure") else \
                ' class="todo"' if "[[" in b[1] else ""
            out.append(f"<p{cls}>{h_inline(b[1])}</p>")
        elif k == "img":
            out.append(f'<figure class="wide"><img alt="{html.escape(b[2])}" src="data:image/png;base64,{b64(b[1])}">'
                       f'<figcaption>{html.escape(b[2])}</figcaption></figure>')
        elif k == "bullets" and section == "In brief":
            cards = []
            for item in b[1]:
                head, _, rest = item.partition("** ")
                cards.append(f'<div class="card"><h3>{html.escape(head.strip("*"))}</h3><p>{h_inline(rest)}</p></div>')
            out.append(f'<div class="cards">{"".join(cards)}</div>')
        elif k == "bullets":
            out.append("<ul>" + "".join(f"<li>{h_inline(t)}</li>" for t in b[1]) + "</ul>")
        elif k == "steps":
            alt = "Flow: " + "; ".join(f"{i + 1}. {h.replace(chr(10), ' ')}" for i, h in enumerate(FLOW))
            out.append(f'<figure><img alt="{html.escape(alt)}" src="data:image/png;base64,{b64(diagram)}"></figure>')
            out.append('<ol class="steps">' + "".join(f"<li><span>{h_inline(t)}</span></li>" for t in b[1]) + "</ol>")
        elif k == "table":
            head = "".join(f"<th>{h_inline(x)}</th>" for x in b[1])
            body = "".join("<tr>" + "".join(f"<td>{h_inline(v)}</td>" for v in r) + "</tr>" for r in b[2])
            out.append(f'<div class="tw"><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></div>')
    page = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>How Mo works</title>
<link rel="preconnect" href="https://fonts.googleapis.com"><link rel="stylesheet"
 href="https://fonts.googleapis.com/css2?family=Inter:wght@400;540;600;700&family=Newsreader:opsz,wght@6..72,470&display=swap">
<style>{CSS}</style></head><body>
<div class="bar"><b>How Mo works · {VERSION}</b>
<button id="edit">Editing: OFF</button><button id="save">Save edited copy</button>
<button id="print" class="primary">Print / Save as PDF</button>
<span class="hint">Turn editing on to change any text or table cell. 'Save edited copy' downloads your version as a new
file; 'Print / Save as PDF' opens your browser's print window.</span></div>
<main>{"".join(out)}<footer>How Mo works · Aevah for BUILT · {ISSUED} · {VERSION}</footer></main>
<script>{JS}</script></body></html>"""
    path.write_text(page, encoding="utf-8")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-pdf", action="store_true")
    a = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    B = parse(SRC.read_text(encoding="utf-8"))
    diagram = flow_diagram(OUT / "mo_flow.png")
    docx_p, pdf_p, html_p = OUT / f"{STEM}.docx", OUT / f"{STEM}.pdf", OUT / f"{STEM}.html"
    render_docx(B, diagram, docx_p)
    render_html(B, diagram, html_p)
    print("wrote", docx_p.relative_to(ROOT), html_p.relative_to(ROOT))
    if not a.no_pdf:
        print("pdf", "ok" if to_pdf(docx_p, pdf_p) else "FAILED", pdf_p.relative_to(ROOT))


if __name__ == "__main__":
    main()
